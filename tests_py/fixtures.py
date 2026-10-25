"""Real documents for the tests. Nothing here is a mock: every function writes an actual file that Reader then opens.

PDFs are drawn with reportlab (known words at known positions, so tests can check rectangles against pixels), SLDC files are
built with the authoritative ``src.utils.sldc``, and the other formats are real DOCX/ODT/EPUB/RTF/HTML/Markdown/TXT files.
"""
from __future__ import annotations

import io
import sys
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
APP_ROOT = HERE.parent


def slai_root() -> Path | None:
    import os
    candidates = [os.environ.get("SLAI_ROOT"), str(APP_ROOT.parent.parent), "/home/claude/slai"]
    for c in candidates:
        if c and (Path(c) / "src" / "utils" / "sldc.py").is_file():
            return Path(c)
    return None


ROOT = slai_root()
if ROOT and str(ROOT) not in sys.path:
    sys.path.append(str(ROOT))
if str(APP_ROOT) not in sys.path:
    sys.path.insert(0, str(APP_ROOT))

LOREM = ("The quick brown fox jumps over the lazy dog. Reading is the act of looking at written symbols and understanding them. "
         "Annotations keep a reader's thoughts attached to the words that caused them. ")


# ---------------------------------------------------------------------------------------------- PDF
def make_pdf(path: Path, pages: int = 3, size: tuple[float, float] = (612, 792), rotate: int = 0, text_prefix: str = "Page",
             sizes: list[tuple[float, float]] | None = None, rotations: list[int] | None = None, bookmarks: bool = True,
             link: bool = True) -> Path:
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(str(path), pagesize=size)
    c.setTitle("Fixture PDF")
    c.setAuthor("SLAI tests")
    for i in range(pages):
        ps = sizes[i] if sizes else size
        c.setPageSize(ps)
        rot = rotations[i] if rotations else rotate
        c.setPageRotation(rot)          # reportlab keeps the previous page's rotation unless it is set again
        w, h = ps
        if rot % 180 == 90:        # reportlab swaps the media box of a rotated page; draw inside it
            w, h = h, w
        c.setFont("Helvetica-Bold", 24)
        c.drawString(72, h - 100, f"{text_prefix} {i + 1} heading")
        c.setFont("Helvetica", 12)
        y = h - 140
        for line in range(12):
            c.drawString(72, y, f"{text_prefix} {i + 1} line {line + 1}: {LOREM[:70]}")
            y -= 18
        c.setFillColorRGB(0.2, 0.5, 0.9)
        c.rect(72, 120, 80, 40, fill=1, stroke=0)
        c.setFillColorRGB(0, 0, 0)
        if bookmarks:
            key = f"p{i + 1}"
            c.bookmarkPage(key)
            c.addOutlineEntry(f"{text_prefix} {i + 1}", key, level=0)
        if link and i == 0 and pages > 1:
            c.linkAbsolute("Go to page 2", "p2", (72, h - 120, 300, h - 100)) if bookmarks else None
        c.showPage()
    c.save()
    return path


def make_scanned_pdf(path: Path, pages: int = 2) -> Path:
    """A PDF whose pages are only images of text: no text layer."""
    from PIL import Image, ImageDraw
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(str(path), pagesize=(612, 792))
    for i in range(pages):
        img = Image.new("RGB", (612, 792), "white")
        d = ImageDraw.Draw(img)
        for line in range(10):
            d.text((60, 80 + line * 24), f"Scanned page {i + 1} line {line + 1} only pixels", fill="black")
        c.drawImage(ImageReader(img), 0, 0, 612, 792)
        c.showPage()
    c.save()
    return path


def make_encrypted_pdf(path: Path, src: Path, user_password: str = "open-sesame", owner_password: str = "owner-secret",
                       allow_copy: bool = True) -> Path:
    import pypdf
    from pypdf.constants import UserAccessPermissions as P
    reader = pypdf.PdfReader(str(src))
    writer = pypdf.PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    perms = P.PRINT | P.ADD_OR_MODIFY | P.MODIFY | P.ADD_OR_MODIFY
    if allow_copy:
        perms |= P.EXTRACT
    writer.encrypt(user_password, owner_password, algorithm="AES-256", permissions_flag=perms)
    with open(path, "wb") as fh:
        writer.write(fh)
    return path


def make_big_pdf(path: Path, pages: int = 400) -> Path:
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(str(path), pagesize=(612, 792))
    for i in range(pages):
        c.setFont("Helvetica", 11)
        for line in range(40):
            c.drawString(60, 740 - line * 16, f"Page {i + 1} line {line + 1} needle{i % 50} {LOREM[:60]}")
        c.showPage()
    c.save()
    return path


# ---------------------------------------------------------------------------------------------- SLDC
def make_sldc(path: Path, paragraphs: int = 30, title: str = "Fixture Document", doc_id: str | None = None, language: str = "en",
              with_table: bool = True, with_figure: bool = True, with_note: bool = True, extra_text: str = "",
              rtl: bool = False) -> Path:
    from src.utils import sldc as s
    blocks = [s.make_block("heading", title, block_id="h1", attributes={"level": 1})]
    refs = []
    assets = []
    for i in range(paragraphs):
        if i % 10 == 0 and i:
            blocks.append(s.make_block("heading", f"Chapter {i // 10 + 1}", block_id=f"h{i // 10 + 1}x", attributes={"level": 2}))
        text = (f"Paragraph {i + 1}. " + LOREM * 2 + extra_text).strip()
        if rtl and i == 1:
            text = "هذا نص عربي للاختبار " * 6
        spans = [{"type": "format", "start": 0, "end": 9, "bold": True}]
        block = s.make_block("paragraph", text, block_id=f"p{i + 1}", spans=spans)
        blocks.append(block)
    if with_note:
        refs.append({"id": "fn1", "kind": "footnote", "text": "A footnote that explains the point."})
        t = blocks[3]["text"]
        blocks[3]["spans"] = sorted(blocks[3].get("spans", []) + [{"type": "ref", "start": len(t) - 5, "end": len(t) - 1, "ref": "fn1"}],
                                    key=lambda sp: (sp["start"], sp["end"], sp["type"]))
    if with_table:
        blocks.append(s.make_table([["Name", "Value"], ["alpha", "1"], ["beta", "2"]], header_rows=1, caption="Table 1. Results", block_id="tbl1"))
    if with_figure:
        from PIL import Image
        buf = io.BytesIO()
        Image.new("RGB", (64, 48), (200, 60, 60)).save(buf, "PNG")
        asset = s.make_asset(buf.getvalue(), asset_id="fig1")
        assets.append(asset)
        blocks.append(s.make_figure(asset, alt="A red rectangle", caption="Figure 1", block_id="fig1b"))
    blocks.append(s.make_block("quote", "A closing quotation worth remembering.", block_id="q1"))
    blocks.append(s.make_block("code", "print('hello')", block_id="c1", attributes={"language": "python"}))
    pkg = s.new_package(title, blocks, language=language, doc_id=doc_id or "fixture-doc-1", references=refs, assets=assets)
    s.write_package(path, pkg, overwrite=True)
    return path


def make_encrypted_sldc(path: Path, password: str = "correct horse battery") -> Path:
    from src.utils import sldc as s
    pkg = s.new_package("Secret", [s.make_block("paragraph", "Top secret reading material.", block_id="p1")], doc_id="secret-doc")
    s.write_package(path, pkg, overwrite=True, recipients=[s.password_recipient(password)])
    return path


def make_signed_sldc(path: Path) -> Path:
    from src.utils import sldc as s
    pkg = s.new_package("Signed", [s.make_block("paragraph", "Signed content that must stay byte-identical.", block_id="p1")], doc_id="signed-doc")
    signer = s.Ed25519Signer.generate()
    pkg = s.sign_package(pkg, signer, purpose="approval")
    s.write_package(path, pkg, overwrite=True)
    return path


def make_commented_sldc(path: Path) -> Path:
    from src.utils import sldc as s
    pkg = s.new_package("Reviewed", [s.make_block("paragraph", "First paragraph to review carefully.", block_id="p1"),
                                     s.make_block("paragraph", "Second paragraph.", block_id="p2")], doc_id="reviewed-doc")
    pkg, _ = s.add_comment(pkg, "p1", "Please check this sentence.", author="reviewer", start=0, end=5)
    s.write_package(path, pkg, overwrite=True)
    return path


# ---------------------------------------------------------------------------------------------- other formats
def make_docx(path: Path) -> Path:
    import docx
    d = docx.Document()
    d.add_heading("DOCX Fixture", 1)
    d.add_paragraph("First paragraph of the Word document. " + LOREM)
    d.add_heading("Second heading", 2)
    d.add_paragraph("Second paragraph with more words to read. " + LOREM)
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "a", "b", "c", "d"
    d.save(str(path))
    return path


def make_odt(path: Path) -> Path:
    content = ('<?xml version="1.0" encoding="UTF-8"?><office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
               'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"><office:body><office:text>'
               '<text:h text:outline-level="1">ODT Fixture</text:h>'
               f'<text:p>First paragraph of the OpenDocument file. {LOREM}</text:p>'
               f'<text:p>Second paragraph. {LOREM}</text:p></office:text></office:body></office:document-content>')
    manifest = ('<?xml version="1.0" encoding="UTF-8"?><manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.2">'
                '<manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.text"/>'
                '<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/></manifest:manifest>')
    with zipfile.ZipFile(path, "w") as z:
        z.writestr(zipfile.ZipInfo("mimetype"), "application/vnd.oasis.opendocument.text")
        z.writestr("content.xml", content)
        z.writestr("META-INF/manifest.xml", manifest)
    return path


def make_epub(path: Path, chapters: int = 3, drm: bool = False) -> Path:
    container = ('<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
                 '<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>')
    items = "".join(f'<item id="c{i}" href="ch{i}.xhtml" media-type="application/xhtml+xml"/>' for i in range(chapters))
    spine = "".join(f'<itemref idref="c{i}"/>' for i in range(chapters))
    opf = ('<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="id">'
           '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>EPUB Fixture</dc:title><dc:creator>Test Author</dc:creator>'
           '<dc:identifier id="id">urn:uuid:12345</dc:identifier><dc:language>en</dc:language></metadata>'
           f'<manifest>{items}</manifest><spine>{spine}</spine></package>')
    with zipfile.ZipFile(path, "w") as z:
        z.writestr(zipfile.ZipInfo("mimetype"), "application/epub+zip")
        z.writestr("META-INF/container.xml", container)
        if drm:
            z.writestr("META-INF/encryption.xml", "<encryption/>")
        z.writestr("OEBPS/content.opf", opf)
        for i in range(chapters):
            z.writestr(f"OEBPS/ch{i}.xhtml", f'<?xml version="1.0"?><html xmlns="http://www.w3.org/1999/xhtml"><head><title>Ch {i}</title></head>'
                                              f'<body><h1>Chapter {i + 1}</h1><p>{LOREM}</p><p>Second paragraph of chapter {i + 1}. {LOREM}</p></body></html>')
    return path


def make_rtf(path: Path) -> Path:
    path.write_text(r"{\rtf1\ansi\deff0 {\fonttbl{\f0 Times New Roman;}}\f0\fs24 RTF fixture paragraph one.\par Second RTF paragraph.\par}", encoding="ascii")
    return path


def make_html(path: Path) -> Path:
    path.write_text(f"<!doctype html><html lang='en'><head><title>HTML Fixture</title><script>alert(1)</script></head>"
                    f"<body><h1>HTML Fixture</h1><p>{LOREM}</p><p>Second paragraph <b>with bold</b> text.</p></body></html>", encoding="utf-8")
    return path


def make_markdown(path: Path) -> Path:
    path.write_text(f"# Markdown Fixture\n\n{LOREM}\n\n## Section\n\n- item one\n- item two\n\nClosing paragraph. {LOREM}\n", encoding="utf-8")
    return path


def make_text(path: Path) -> Path:
    path.write_text("Plain text fixture.\n\n" + (LOREM + "\n\n") * 3, encoding="utf-8")
    return path


# ---------------------------------------------------------------------------------------------- damaged files
def make_truncated_pdf(path: Path, src: Path) -> Path:
    data = src.read_bytes()
    path.write_bytes(data[: len(data) // 3])
    return path


def make_garbage(path: Path, magic: bytes = b"PK\x03\x04") -> Path:
    path.write_bytes(magic + b"\x00garbage-that-is-not-a-document" * 20)
    return path
